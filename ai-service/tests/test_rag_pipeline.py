from io import BytesIO

from docx import Document

from rag_pipeline import PipelineError, chunk_text, clean_text, extract_document


def text_pdf(value: str) -> bytes:
    stream = f'BT /F1 18 Tf 72 72 Td ({value}) Tj ET'
    objects = [
        '<< /Type /Catalog /Pages 2 0 R >>',
        '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>',
        f'<< /Length {len(stream)} >>\nstream\n{stream}\nendstream',
        '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    data = b'%PDF-1.4\n'
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(data))
        data += f'{index} 0 obj\n{body}\nendobj\n'.encode()
    xref = len(data)
    data += f'xref\n0 {len(objects) + 1}\n0000000000 65535 f \n'.encode()
    data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
    return data + f'trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()


def test_clean_text_normalizes_whitespace() -> None:
    assert clean_text('  Xin\tchào\r\n\n\nFLTS  ') == 'Xin chào\n\nFLTS'


def test_chunk_text_is_deterministic_and_keeps_overlap() -> None:
    text = '\n\n'.join([' '.join(f'word{index}_{item}' for item in range(90)) for index in range(4)])
    chunks = chunk_text(text, target_words=120, overlap_words=10)
    assert len(chunks) == 4
    assert chunks[0]['chunk_index'] == 0
    assert chunks[0]['content_hash']
    assert 'word0_89' in chunks[1]['content']


def test_legacy_doc_uses_the_selected_parser(monkeypatch) -> None:
    class Result:
        returncode = 0
        stdout = 'Legacy FLTS document'

    monkeypatch.setattr('rag_pipeline.subprocess.run', lambda *args, **kwargs: Result())
    result = extract_document(b'binary-doc', 'doc')
    assert result.text == 'Legacy FLTS document'
    assert result.metadata['parser'] == 'antiword'


def test_legacy_doc_error_is_reported_honestly(monkeypatch) -> None:
    class Result:
        returncode = 1
        stdout = ''

    monkeypatch.setattr('rag_pipeline.subprocess.run', lambda *args, **kwargs: Result())
    try:
        extract_document(b'not a real word document', 'doc')
    except PipelineError as error:
        assert error.code == 'DOC_EXTRACTION_FAILED'
    else:
        raise AssertionError('Legacy DOC must not be presented as successfully extracted.')


def test_extracts_text_based_pdf_with_page_metadata() -> None:
    result = extract_document(text_pdf('Hello FLTS'), 'pdf')
    assert result.page_count == 1
    assert result.metadata['parser'] == 'pypdf'
    assert 'Hello FLTS' in result.text


def test_empty_pdf_and_corrupt_pdf_fail_without_false_text() -> None:
    for content in (text_pdf(''), b'not a pdf'):
        try:
            extract_document(content, 'pdf')
        except PipelineError as error:
            assert error.code in {'NO_EXTRACTABLE_TEXT', 'PDF_EXTRACTION_FAILED'}
        else:
            raise AssertionError('An empty or corrupt PDF must not be accepted.')


def test_extracts_docx_paragraphs_and_tables() -> None:
    document = Document()
    document.add_paragraph('FLTS document paragraph')
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'Topic'
    table.cell(0, 1).text = 'RAG'
    buffer = BytesIO()
    document.save(buffer)

    result = extract_document(buffer.getvalue(), 'docx')
    assert 'FLTS document paragraph' in result.text
    assert 'Topic | RAG' in result.text
    assert '[Paragraph 1]' in result.text
    assert result.metadata['table_count'] == 1
